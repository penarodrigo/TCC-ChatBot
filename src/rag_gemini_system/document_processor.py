# -*- coding: utf-8 -*-
"""
Módulo para Processamento e Extração de Texto de Documentos.
"""

import logging
from pathlib import Path
from typing import Union

# Importações para processamento de documentos
import pandas as pd
import pdfplumber
from pptx import Presentation, exc as pptx_exc
import docx
from docx.opc import exceptions as docx_exc
from openpyxl import load_workbook
from openpyxl.utils import exceptions as openpyxl_exc

logger = logging.getLogger(__name__)

class DocumentProcessor:
    """
    Classe estática para extrair texto de vários tipos de arquivo.
    """
    SUPPORTED_EXTENSIONS = {'.txt', '.pdf', '.pptx', '.csv', '.docx', '.xlsx'}

    @classmethod
    def is_supported(cls, file_path: Union[str, Path]) -> bool:
        """Verifica se a extensão do arquivo é suportada."""
        return Path(file_path).suffix.lower() in cls.SUPPORTED_EXTENSIONS

    @staticmethod
    def read_txt(path: Path) -> str:
        """Lê um arquivo de texto com tratamento para diferentes encodings."""
        logger.debug(f"Lendo arquivo TXT: {path}")
        try:
            return path.read_text(encoding='utf-8')
        except UnicodeDecodeError:
            logger.warning(f"Falha ao decodificar {path} como UTF-8. Tentando com Latin-1.")
            try:
                return path.read_text(encoding='latin-1')
            except Exception as e:
                logger.error(f"Não foi possível ler o arquivo {path} com Latin-1: {e}")
                return ""
        except Exception as e:
            logger.error(f"Erro ao ler o arquivo TXT {path}: {e}")
            return ""

    @staticmethod
    def read_pdf(path: Path) -> str:
        """Extrai texto de um arquivo PDF."""
        logger.debug(f"Lendo arquivo PDF: {path}")
        text = []
        try:
            with pdfplumber.open(path) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text.append(page_text)
            return "\n".join(text)
        except pdfplumber.exceptions.PDFSyntaxError as e:
            logger.error(f"Erro de sintaxe no PDF {path}: {e}")
            return ""
        except Exception as e:
            logger.error(f"Erro ao ler o arquivo PDF {path}: {e}")
            return ""

    @staticmethod
    def read_docx(path: Path) -> str:
        """Extrai texto de um arquivo DOCX."""
        logger.debug(f"Lendo arquivo DOCX: {path}")
        try:
            doc = docx.Document(path)
            return "\n".join([p.text for p in doc.paragraphs])
        except docx_exc.PackageNotFoundError:
            logger.error(f"Arquivo DOCX não encontrado ou corrompido: {path}")
            return ""
        except Exception as e:
            logger.error(f"Erro ao ler o arquivo DOCX {path}: {e}")
            return ""

    @staticmethod
    def read_pptx(path: Path) -> str:
        """Extrai texto de um arquivo PPTX."""
        logger.debug(f"Lendo arquivo PPTX: {path}")
        text = []
        try:
            prs = Presentation(path)
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text.append(shape.text)
            return "\n".join(text)
        except pptx_exc.PackageNotFoundError:
            logger.error(f"Arquivo PPTX não encontrado ou corrompido: {path}")
            return ""
        except Exception as e:
            logger.error(f"Erro ao ler o arquivo PPTX {path}: {e}")
            return ""

    @staticmethod
    def read_csv(path: Path) -> str:
        """Lê dados de um arquivo CSV com tratamento de erros."""
        logger.debug(f"Lendo arquivo CSV: {path}")
        try:
            df = pd.read_csv(path, encoding='utf-8')
        except (UnicodeDecodeError, pd.errors.ParserError):
            logger.warning(f"Falha no parse de {path} com UTF-8. Tentando alternativas.")
            try:
                # Tenta ler com outro encoding e sem delimitador definido
                df = pd.read_csv(path, encoding='latin-1', sep=None, engine='python')
            except Exception as inner_e:
                logger.error(f"Falha final ao tentar ler CSV {path}: {inner_e}")
                return ""
        except Exception as e:
            logger.error(f"Erro inesperado ao ler o arquivo CSV {path}: {e}")
            return ""
        return df.to_string(index=False)

    @staticmethod
    def read_xlsx(path: Path) -> str:
        """Extrai texto de todas as planilhas de um arquivo XLSX."""
        logger.debug(f"Lendo arquivo XLSX: {path}")
        text_content = []
        try:
            wb = load_workbook(filename=path, read_only=True, data_only=True)
            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                text_content.append(f"Planilha: {sheet_name}")
                for row in sheet.iter_rows():
                    row_values = [str(c.value) if c.value is not None else "" for c in row]
                    text_content.append(", ".join(row_values))
                text_content.append("")  # Linha em branco entre planilhas
            return "\n".join(text_content)
        except openpyxl_exc.InvalidFileException:
            logger.error(f"Arquivo XLSX inválido ou não encontrado: {path}")
            return ""
        except Exception as e:
            logger.error(f"Erro ao ler o arquivo XLSX {path}: {e}")
            return ""

    @classmethod
    def extract_text(cls, file_path: Union[str, Path]) -> str:
        """
        Extrai texto de um arquivo, verificando o tipo e tratando exceções.
        """
        path_obj = Path(file_path)
        ext = path_obj.suffix.lower()
        logger.info(f"Processando arquivo: {file_path} (tipo: {ext})")

        if not path_obj.is_file():
            logger.error(f"Arquivo não encontrado: {file_path}")
            return ""

        if not cls.is_supported(path_obj):
            logger.warning(f"Tipo de arquivo não suportado: {ext} para {file_path}")
            return ""

        readers = {
            '.txt': cls.read_txt,
            '.pdf': cls.read_pdf,
            '.docx': cls.read_docx,
            '.pptx': cls.read_pptx,
            '.csv': cls.read_csv,
            '.xlsx': cls.read_xlsx,
        }

        reader = readers.get(ext)
        
        # A verificação is_supported() garante que o reader não será None
        try:
            content = reader(path_obj)
            if not content or not content.strip():
                logger.warning(f"Nenhum conteúdo extraído de '{file_path}'. O arquivo pode estar vazio.")
                return ""
            return content.strip()
        except FileNotFoundError:
             logger.error(f"Arquivo não encontrado durante a leitura: {file_path}")
             return ""
        except Exception as e:
            logger.error(f"Erro inesperado ao extrair texto de {file_path}: {e}", exc_info=True)
            return ""
